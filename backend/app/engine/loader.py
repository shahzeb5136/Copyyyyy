"""Read an uploaded CSV/Excel file into a typed DataFrame.

Responsibilities
* encoding + delimiter sniffing for CSV/TSV/TXT, openpyxl for .xlsx/.xlsm
* header validation: the FIRST row must contain the column names (clear error otherwise)
* per-column type inference (integer, number, date, datetime, boolean, text) and coercion
* row/column limits
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

from ..errors import FileFormatError

DATA_TYPES = ["number", "integer", "date", "datetime", "boolean", "text"]
CSV_EXTENSIONS = {".csv", ".tsv", ".txt"}
EXCEL_EXTENSIONS = {".xlsx", ".xlsm"}
SUPPORTED_EXTENSIONS = CSV_EXTENSIONS | EXCEL_EXTENSIONS

HEADER_MESSAGE = ("The first row of the file must contain the column names, followed by the data rows. "
                  "Please put the column names in the first row (remove any title rows, notes or blank rows "
                  "above the header) and upload the file again.")

_TRUE = {"true", "t", "yes", "y", "1"}
_FALSE = {"false", "f", "no", "n", "0"}
_NUMERIC_CLEAN = re.compile(r"[,\s$€£¥%]|AED|SAR|USD|EUR|GBP", re.IGNORECASE)
_DATE_LIKE = re.compile(r"\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}|\d{1,2}\s+[A-Za-z]{3,}\.?\s+\d{2,4}|[A-Za-z]{3,}\s+\d{1,2},?\s+\d{4}")


@dataclass
class LoadedTable:
    df: pd.DataFrame
    types: dict[str, str]
    sheet_names: list[str] = field(default_factory=list)
    sheet: str = ""
    warnings: list[str] = field(default_factory=list)
    coercion_failures: dict[str, int] = field(default_factory=dict)

    @property
    def row_count(self) -> int:
        return len(self.df)


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #
def file_kind(filename: str) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in CSV_EXTENSIONS:
        return "csv"
    if ext in EXCEL_EXTENSIONS:
        return "excel"
    if ext == ".xls":
        raise FileFormatError("Legacy .xls files are not supported. Please save the workbook as .xlsx and upload again.")
    raise FileFormatError(f"Unsupported file type '{ext or '(none)'}'. Upload a CSV (.csv/.tsv/.txt) or Excel (.xlsx) file.")


def list_sheets(content: bytes) -> list[str]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise FileFormatError(f"The Excel file could not be opened ({type(exc).__name__}). Is it a valid .xlsx file?") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_table(content: bytes, filename: str, sheet: str | None = None, max_rows: int = 1_000_000,
               max_columns: int = 300) -> LoadedTable:
    kind = file_kind(filename)
    if not content:
        raise FileFormatError("The uploaded file is empty.")
    warnings: list[str] = []
    sheet_names: list[str] = []
    if kind == "csv":
        raw = _read_csv(content, max_rows)
        sheet_used = ""
    else:
        sheet_names = list_sheets(content)
        if not sheet_names:
            raise FileFormatError("The workbook has no sheets.")
        sheet_used = sheet or sheet_names[0]
        if sheet_used not in sheet_names:
            raise FileFormatError(f"Sheet '{sheet_used}' not found. Available sheets: {', '.join(sheet_names)}")
        raw = _read_excel(content, sheet_used, max_rows)
        if len(sheet_names) > 1 and not sheet:
            warnings.append(f"The workbook has {len(sheet_names)} sheets; using '{sheet_used}'. "
                            "Pick another sheet if this is not the data you want.")

    raw = raw.dropna(how="all").reset_index(drop=True)
    if raw.empty or raw.shape[1] == 0:
        raise FileFormatError("The file contains no data rows.")
    if raw.shape[1] > max_columns:
        raise FileFormatError(f"The file has {raw.shape[1]} columns; the limit is {max_columns}.")

    header, body = _split_header(raw, warnings)
    if len(body) > max_rows:
        raise FileFormatError(f"The file has more than {max_rows:,} data rows. Please aggregate or split the data.")
    if body.empty:
        raise FileFormatError("The file has column names but no data rows.")

    body.columns = header
    # Drop columns that are entirely empty and had an empty header.
    drop = [c for c in body.columns if c.startswith("Column_") and body[c].isna().all()]
    if drop:
        body = body.drop(columns=drop)
    body = body.dropna(how="all").reset_index(drop=True)

    types: dict[str, str] = {}
    failures: dict[str, int] = {}
    out = pd.DataFrame(index=body.index)
    for col in body.columns:
        t = infer_type(body[col])
        coerced, failed = coerce_series(body[col], t)
        if failed and failed > 0.05 * max(1, coerced.notna().sum() + failed):
            # Too many values don't fit the guessed type: keep the column as text instead of losing data.
            coerced, failed, t = coerce_series(body[col], "text")[0], 0, "text"
        out[col] = coerced
        types[col] = t
        if failed:
            failures[col] = failed
            warnings.append(f"Column '{col}': {failed} value(s) could not be read as {t} and were left blank.")
    return LoadedTable(df=out, types=types, sheet_names=sheet_names, sheet=sheet_used, warnings=warnings,
                       coercion_failures=failures)


# --------------------------------------------------------------------------- #
# Readers (header=None: we validate the header row ourselves)
# --------------------------------------------------------------------------- #
def _read_csv(content: bytes, max_rows: int) -> pd.DataFrame:
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None or not text.strip():
        raise FileFormatError("The file is empty or is not a text/CSV file.")
    sample = text[:50_000]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    _check_first_line_width(sample, delimiter)
    try:
        df = pd.read_csv(io.StringIO(text), sep=delimiter, header=None, dtype=str, keep_default_na=True,
                         nrows=max_rows + 2, skipinitialspace=True, on_bad_lines="skip",
                         engine="python" if delimiter == "|" else "c")
    except Exception as exc:  # noqa: BLE001
        raise FileFormatError(f"The CSV file could not be parsed ({type(exc).__name__}: {str(exc)[:200]}).") from exc
    return df


def _check_first_line_width(sample: str, delimiter: str) -> None:
    """A title row ('Budget report 2026') has fewer fields than the data rows below it."""
    rows = [r for r in csv.reader(io.StringIO(sample), delimiter=delimiter) if any(c.strip() for c in r)]
    if len(rows) < 3:
        return
    widths = [len(r) for r in rows[1:25]]
    typical = max(set(widths), key=widths.count)
    if len(rows[0]) < typical and typical >= 2 and len(rows[0]) <= max(1, typical // 2):
        raise FileFormatError(f"The first row has {len(rows[0])} field(s) while the data rows have {typical}, so it looks "
                              "like a title row rather than the column names. " + HEADER_MESSAGE)


def _read_excel(content: bytes, sheet: str, max_rows: int) -> pd.DataFrame:
    try:
        df = pd.read_excel(io.BytesIO(content), sheet_name=sheet, header=None, engine="openpyxl",
                           nrows=max_rows + 2)
    except Exception as exc:  # noqa: BLE001
        raise FileFormatError(f"Sheet '{sheet}' could not be read ({type(exc).__name__}: {str(exc)[:200]}).") from exc
    return df


# --------------------------------------------------------------------------- #
# Header validation
# --------------------------------------------------------------------------- #
def _cell_is_blank(v) -> bool:
    return v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and not v.strip()) or v is pd.NaT


def _looks_like_data(v) -> bool:
    """True when a header cell looks like a value (number/date/bool) rather than a column name."""
    if isinstance(v, (bool, np.bool_)):
        return True
    if isinstance(v, (int, float, np.integer, np.floating)) and not _cell_is_blank(v):
        return True
    if isinstance(v, (datetime, date, pd.Timestamp)):
        return True
    s = str(v).strip()
    if not s:
        return False
    if re.fullmatch(r"[-+]?[\d,.]+%?", s) and any(ch.isdigit() for ch in s):
        return True
    if _DATE_LIKE.fullmatch(s):
        return True
    return s.lower() in _TRUE | _FALSE and s.lower() not in {"y", "n", "t", "f"}


def _split_header(raw: pd.DataFrame, warnings: list[str]) -> tuple[list[str], pd.DataFrame]:
    first = list(raw.iloc[0])
    n_cols = len(first)
    blank = [_cell_is_blank(v) for v in first]
    n_blank = sum(blank)
    non_blank = [v for v, b in zip(first, blank) if not b]
    body = raw.iloc[1:].reset_index(drop=True)
    # Which columns actually hold data below the header?
    has_data = [bool(body[c].notna().any()) for c in raw.columns]
    blank_with_data = sum(1 for b, d in zip(blank, has_data) if b and d)

    if not non_blank:
        raise FileFormatError("The first row is empty. " + HEADER_MESSAGE)
    if n_cols >= 3 and len(non_blank) == 1 and sum(has_data) >= 2:
        raise FileFormatError(f"The first row contains a single cell ('{str(non_blank[0])[:60]}'), which looks like "
                              "a title rather than column names. " + HEADER_MESSAGE)
    data_like = sum(1 for v in non_blank if _looks_like_data(v))
    if data_like >= max(1, len(non_blank) * 0.5):
        raise FileFormatError("The first row looks like a data row (numbers or dates) rather than column names. "
                              + HEADER_MESSAGE)
    if blank_with_data and blank_with_data / max(1, sum(has_data)) > 0.3:
        raise FileFormatError(f"{blank_with_data} column(s) with data have no name in the first row. " + HEADER_MESSAGE)
    if n_blank:
        warnings.append(f"{n_blank} header cell(s) were empty; those columns were named Column_<n> or dropped if empty.")

    names: list[str] = []
    seen: dict[str, int] = {}
    for i, v in enumerate(first):
        name = f"Column_{i + 1}" if blank[i] else _clean_header(v)
        base = name
        if base in seen:
            seen[base] += 1
            name = f"{base}_{seen[base]}"
            warnings.append(f"Duplicate column name '{base}' renamed to '{name}'.")
        else:
            seen[base] = 1
        names.append(name)
    return names, body


def _clean_header(v) -> str:
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = " ".join(str(v).replace("\n", " ").split())
    return s[:120] or "Column"


# --------------------------------------------------------------------------- #
# Type inference and coercion
# --------------------------------------------------------------------------- #
def _is_stringy(s: pd.Series) -> bool:
    return s.dtype == object or pd.api.types.is_string_dtype(s)


def infer_type(s: pd.Series) -> str:
    """Best-effort type guess for a raw column (object/str from CSV, typed from Excel)."""
    non_null = s.dropna()
    if non_null.empty:
        return "text"
    if pd.api.types.is_bool_dtype(s):
        return "boolean"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "date" if (non_null.dt.normalize() == non_null).all() else "datetime"
    if pd.api.types.is_integer_dtype(s):
        return "integer"
    if pd.api.types.is_float_dtype(s):
        return "integer" if (non_null.round() == non_null).all() and non_null.abs().max() < 2**53 else "number"

    # object column (Excel mixed cells) - look at python types first
    if s.dtype == object:
        kinds = {type(v) for v in non_null.head(5_000)}
        if kinds <= {bool, np.bool_}:
            return "boolean"
        if kinds <= {int, float, np.integer, np.floating, bool, np.bool_} and bool not in kinds:
            vals = pd.to_numeric(non_null, errors="coerce")
            return "integer" if (vals.dropna().round() == vals.dropna()).all() else "number"
        if kinds <= {datetime, date, pd.Timestamp}:
            parsed = pd.to_datetime(non_null, errors="coerce")
            return "date" if (parsed.dropna().dt.normalize() == parsed.dropna()).all() else "datetime"

    sample = non_null.astype(str).str.strip()
    sample = sample[sample != ""].head(5_000)
    if sample.empty:
        return "text"
    lowered = sample.str.lower()
    if lowered.isin(_TRUE | _FALSE).all() and lowered.nunique() <= 2 and not lowered.isin({"0", "1"}).all():
        return "boolean"
    numeric = pd.to_numeric(sample.str.replace(_NUMERIC_CLEAN, "", regex=True)
                            .str.replace(r"^\((.*)\)$", r"-\1", regex=True), errors="coerce")
    if numeric.notna().mean() > 0.97:
        return "integer" if (numeric.dropna().round() == numeric.dropna()).all() else "number"
    if sample.str.contains(_DATE_LIKE, regex=True).mean() > 0.9:
        parsed = _parse_dates(sample)
        if parsed.notna().mean() > 0.95:
            has_time = (parsed.dropna().dt.normalize() != parsed.dropna()).any()
            return "datetime" if has_time else "date"
    return "text"


def _detect_dayfirst(s: pd.Series) -> list[bool]:
    parts = s.dropna().astype(str).str.extract(r"^\s*(\d{1,2})[/.\-](\d{1,2})[/.\-]\d{2,4}")
    parts = parts.dropna().astype(int)
    if parts.empty:
        return [False, True]
    if (parts[0] > 12).any() and not (parts[1] > 12).any():
        return [True]
    if (parts[1] > 12).any() and not (parts[0] > 12).any():
        return [False]
    return [False, True]


def _parse_dates(s: pd.Series) -> pd.Series:
    best = pd.to_datetime(s, errors="coerce", format="ISO8601")
    if best.notna().mean() > 0.95:
        return best
    for dayfirst in _detect_dayfirst(s):
        try:
            cand = pd.to_datetime(s, errors="coerce", dayfirst=dayfirst, format="mixed")
        except (ValueError, TypeError):
            continue
        if cand.notna().sum() > best.notna().sum():
            best = cand
    return best


def coerce_series(s: pd.Series, col_type: str) -> tuple[pd.Series, int]:
    """Convert a raw series to the declared type. Returns (series, number of values that failed)."""
    if _is_stringy(s):
        stripped = s.astype("string").str.strip()
        non_null_before = stripped.notna() & (stripped != "")
    else:
        non_null_before = s.notna()
    if col_type == "text":
        out = s.astype("string").str.strip()
        out = out.where(out != "", pd.NA)
        return out, 0
    if col_type in ("integer", "number"):
        raw = s
        if _is_stringy(raw):
            raw = (raw.astype("string").str.strip().str.replace(_NUMERIC_CLEAN, "", regex=True)
                   .str.replace(r"^\((.*)\)$", r"-\1", regex=True))
        out = pd.to_numeric(raw, errors="coerce")
        if col_type == "integer":
            fractional = out.notna() & (out.round() != out)
            out = out.astype("float64") if fractional.any() else out.round().astype("Int64")
        else:
            out = out.astype("float64")
    elif col_type in ("date", "datetime"):
        if pd.api.types.is_datetime64_any_dtype(s):
            out = s
        elif s.dtype == object and all(isinstance(v, (datetime, date, pd.Timestamp)) for v in s.dropna().head(200)):
            out = pd.to_datetime(s, errors="coerce")
        else:
            out = _parse_dates(s.astype("string").str.strip())
        if getattr(out.dt, "tz", None) is not None:
            out = out.dt.tz_localize(None)
        if col_type == "date":
            out = out.dt.normalize()
    elif col_type == "boolean":
        lowered = s.astype("string").str.strip().str.lower()
        out = pd.Series(pd.NA, index=s.index, dtype="boolean")
        out[lowered.isin(_TRUE)] = True
        out[lowered.isin(_FALSE)] = False
    else:
        raise ValueError(f"Unknown column type: {col_type}")
    failed = int((non_null_before & out.isna()).sum())
    return out, failed


def coerce_frame(df: pd.DataFrame, types: dict[str, str]) -> tuple[pd.DataFrame, dict[str, int]]:
    out = pd.DataFrame(index=df.index)
    failures: dict[str, int] = {}
    for col in df.columns:
        out[col], failed = coerce_series(df[col], types.get(col, "text"))
        if failed:
            failures[col] = failed
    return out, failures
